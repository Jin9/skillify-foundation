#!/usr/bin/env python3
"""Convert a CSV file into an XLSX workbook with an AutoFilter on the header row.

Fidelity-first: every cell is written as text (number format '@'), so long digit
IDs never collapse into scientific notation and leading zeros / timestamp seconds
survive. Values that look like formulas are neutralized to literal text. Two-pass
streaming design: the input is fully validated before any output file is created.

stdout on success (single final line):
  OK input=... output=... sheet=... rows=N cols=N filter=A1:CN mode=minimal|styled warnings=N
stderr: "WARN ..." (non-fatal) and "ERROR ..." (fatal) lines.
Exit codes: 0 success, 1 input/validation/environment error, 2 usage error.
"""

import argparse
import codecs
import csv
import re
import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.exceptions import IllegalCharacterError
except ImportError:
    sys.stderr.write(
        "ERROR openpyxl is required. Install: python3 -m pip install --user openpyxl"
        " (or: uv pip install openpyxl)\n"
    )
    sys.exit(1)

EXCEL_MAX_ROWS = 1048576
EXCEL_MAX_COLS = 16384
SHEET_FORBIDDEN = "[]:*?/\\"
SCI_NOTATION = re.compile(r"\d\.\d+E\+\d+")
SCI_SCAN_MAX_LEN = 32  # display-form damage is short; skips long payload cells
WIDTH_SAMPLE_ROWS = 200
WIDTH_MIN = 10
WIDTH_MAX = 60
HEADER_FILL_RGB = "FF1F4E78"
HEADER_FONT_RGB = "FFFFFFFF"
HEADER_ROW_HEIGHT = 26


def die(msg):
    sys.stderr.write("ERROR " + msg + "\n")
    sys.exit(1)


def warn(msg):
    sys.stderr.write("WARN " + msg + "\n")


def raise_field_limit():
    try:
        csv.field_size_limit(sys.maxsize)
    except OverflowError:
        csv.field_size_limit(2 ** 31 - 1)


def sanitize_sheet_name(raw):
    name = "".join(ch for ch in raw if ch not in SHEET_FORBIDDEN)
    name = name.strip()[:31].strip()
    return name or "Sheet1"


def open_csv(path, encoding):
    return open(str(path), "r", newline="", encoding=encoding)


def scan(path, delimiter, encoding):
    """Pass 1: validate and measure without creating any output."""
    header = None
    n_rows = 0
    short_rows = 0
    maxlens = []
    sci_count = 0
    sci_first = None
    sci_sample = None
    with open_csv(path, encoding) as fh:
        reader = csv.reader(fh, delimiter=delimiter)
        for rec, row in enumerate(reader, start=1):  # rec == spreadsheet row number
            if header is None:
                header = row
                if not header or all(v.strip() == "" for v in header):
                    die("input has an empty header row: %s" % path)
                if len(header) > EXCEL_MAX_COLS:
                    die("%d columns exceed the XLSX limit of %d" % (len(header), EXCEL_MAX_COLS))
                maxlens = [len(v) for v in header]
                continue
            n_rows += 1
            if len(row) > len(header):
                die(
                    "row %d has %d fields, expected %d (delimiter or quoting mismatch?)"
                    % (rec, len(row), len(header))
                )
            if len(row) < len(header):
                short_rows += 1
            sample_widths = n_rows <= WIDTH_SAMPLE_ROWS
            for col, value in enumerate(row):
                if sample_widths and len(value) > maxlens[col]:
                    maxlens[col] = len(value)
                if len(value) <= SCI_SCAN_MAX_LEN and SCI_NOTATION.search(value):
                    sci_count += 1
                    if sci_first is None:
                        sci_first = "R%dC%d" % (rec, col + 1)
                        sci_sample = value
    if header is None:
        die("input is empty: %s" % path)
    if n_rows + 1 > EXCEL_MAX_ROWS:
        die("%d data rows exceed the XLSX limit of %d" % (n_rows, EXCEL_MAX_ROWS - 1))
    return header, n_rows, short_rows, maxlens, sci_count, sci_first, sci_sample


def write_xlsx(csv_path, out_path, sheet, styled, delimiter, encoding, header, n_rows, maxlens):
    n_cols = len(header)
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet(title=sheet)
    if styled:
        # views/cols/row-heights flush on the first append - must all be set before it
        for idx in range(1, n_cols + 1):
            width = min(WIDTH_MAX, max(WIDTH_MIN, maxlens[idx - 1] + 2))
            ws.column_dimensions[get_column_letter(idx)].width = width
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = HEADER_ROW_HEIGHT

    header_font = Font(name="Calibri", size=11, bold=True, color=HEADER_FONT_RGB)
    header_fill = PatternFill(fill_type="solid", fgColor=HEADER_FILL_RGB)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def text_cell(value):
        cell = WriteOnlyCell(ws, value=value)
        if cell.data_type == "f":  # "=..." stays literal text, never a live formula
            cell.data_type = "s"
        cell.number_format = "@"
        return cell

    def header_cell(value):
        cell = text_cell(value)
        if styled:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
        return cell

    with open_csv(csv_path, encoding) as fh:
        reader = csv.reader(fh, delimiter=delimiter)
        try:
            first = True
            for row in reader:
                if first:
                    ws.append([header_cell(v) for v in row])
                    first = False
                    continue
                if len(row) < n_cols:
                    row = row + [""] * (n_cols - len(row))
                ws.append([text_cell(v) for v in row])
        except IllegalCharacterError:
            die("input contains control characters that are illegal in XLSX")
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(n_cols), n_rows + 1)
    wb.save(str(out_path))


def main():
    parser = argparse.ArgumentParser(
        description="Convert a CSV into an XLSX with an AutoFilter on the header row,"
        " every value preserved as text."
    )
    parser.add_argument("input", help="source CSV file")
    parser.add_argument("-o", "--output", help="output XLSX path (default: input name with .xlsx)")
    parser.add_argument(
        "--styled",
        action="store_true",
        help="frozen header row, bold white-on-navy header, sized columns",
    )
    parser.add_argument("--sheet", help="worksheet name (default: input stem, sanitized, max 31 chars)")
    parser.add_argument("--delimiter", default=",", help="field delimiter (default ','; use \\t for tab)")
    parser.add_argument("--encoding", default="utf-8-sig", help="input encoding (default utf-8-sig)")
    parser.add_argument("--force", action="store_true", help="overwrite an existing output file")
    args = parser.parse_args()

    raise_field_limit()
    delimiter = "\t" if args.delimiter == "\\t" else args.delimiter
    if len(delimiter) != 1:
        die("--delimiter must be a single character")
    try:
        codecs.lookup(args.encoding)
    except LookupError:
        die("unknown --encoding: %s" % args.encoding)

    csv_path = Path(args.input).resolve()
    if not csv_path.is_file():
        die("input not found: %s" % csv_path)
    out_path = Path(args.output).resolve() if args.output else csv_path.with_suffix(".xlsx")
    if out_path == csv_path:
        die("output path equals input path")
    if not out_path.parent.is_dir():
        die("output directory does not exist: %s" % out_path.parent)
    lock = out_path.parent / ("~$" + out_path.name)
    if lock.exists():
        die("Excel holds a lock on the output (%s); close the workbook first" % lock.name)
    if out_path.exists() and not args.force:
        die("output exists: %s (pass --force to overwrite)" % out_path)
    sheet = sanitize_sheet_name(args.sheet if args.sheet else csv_path.stem)

    try:
        header, n_rows, short_rows, maxlens, sci_count, sci_first, sci_sample = scan(
            csv_path, delimiter, args.encoding
        )
        warnings = 0
        if short_rows:
            warnings += 1
            warn("short_rows=%d padded to %d columns" % (short_rows, len(header)))
        if sci_count:
            warnings += 1
            warn(
                "sci_notation cells=%d first=%s sample=%s"
                " (input already Excel-damaged; conversion stays faithful to the CSV as-is)"
                % (sci_count, sci_first, sci_sample)
            )
        write_xlsx(
            csv_path, out_path, sheet, args.styled, delimiter, args.encoding,
            header, n_rows, maxlens,
        )
    except UnicodeDecodeError as exc:
        die(
            "cannot decode input as %s (%s); try --encoding cp874 or tis-620"
            " for legacy Thai exports" % (args.encoding, exc)
        )
    except csv.Error as exc:
        die("CSV parse error: %s" % exc)

    print(
        "OK input=%s output=%s sheet=%s rows=%d cols=%d filter=A1:%s%d mode=%s warnings=%d"
        % (
            csv_path,
            out_path,
            sheet,
            n_rows,
            len(header),
            get_column_letter(len(header)),
            n_rows + 1,
            "styled" if args.styled else "minimal",
            warnings,
        )
    )


if __name__ == "__main__":
    main()
